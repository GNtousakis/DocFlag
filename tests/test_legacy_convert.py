""".doc conversion tests. Neither Word automation nor LibreOffice is available
where these run, so both are exercised through stand-ins: a real subprocess
that follows soffice's command-line contract, and a fake win32com/pythoncom
that records the calls Word would receive."""
from __future__ import annotations

import io
import os
import sys
import types

import pytest
from docx import Document
from nicegui.elements.upload import Upload
from nicegui.testing import User

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docflag import legacy_convert
from docflag.legacy_convert import convert_doc, is_legacy_doc
from tests.fixtures.make_fixtures import build_plain

_FAKE_SOFFICE_SCRIPT = """
import shutil, sys, os
args = sys.argv[1:]
assert args[:3] == ["--headless", "--convert-to", "docx"], args
outdir = args[args.index("--outdir") + 1]
stem = os.path.splitext(os.path.basename(args[-1]))[0]
shutil.copyfile(os.environ["FAKE_SOFFICE_SOURCE_DOCX"], os.path.join(outdir, stem + ".docx"))
"""


@pytest.fixture
def fake_soffice(tmp_path, monkeypatch):
    """Route conversions to a script that 'converts' any file into a known
    .docx, through the real subprocess call."""
    script = tmp_path / "fake_soffice.py"
    script.write_text(_FAKE_SOFFICE_SCRIPT)
    source = build_plain(str(tmp_path / "converted" / "source.docx"), "An old draft.")
    monkeypatch.setenv("FAKE_SOFFICE_SOURCE_DOCX", source)
    monkeypatch.setattr(legacy_convert.sys, "platform", "linux")
    monkeypatch.setattr(legacy_convert, "_find_soffice", lambda: "soffice")
    real_run = legacy_convert.subprocess.run
    monkeypatch.setattr(
        legacy_convert.subprocess, "run",
        lambda cmd, **kw: real_run([sys.executable, str(script), *cmd[1:]], **kw),
    )


def test_is_legacy_doc():
    assert is_legacy_doc("report.doc") and is_legacy_doc("REPORT.DOC")
    assert not is_legacy_doc("report.docx")


def test_libreoffice_conversion(tmp_path, fake_soffice):
    doc_path = tmp_path / "in_1.doc"
    doc_path.write_bytes(b"stand-in for Word 97-2003 content")
    out = convert_doc(str(doc_path))
    assert out == str(tmp_path / "in_1.docx")
    assert Document(out).paragraphs[0].text == "An old draft."


def test_clear_error_when_no_converter(tmp_path, monkeypatch):
    monkeypatch.setattr(legacy_convert.sys, "platform", "linux")
    monkeypatch.setattr(legacy_convert, "_find_soffice", lambda: None)
    with pytest.raises(RuntimeError, match="Microsoft Word or LibreOffice"):
        convert_doc(str(tmp_path / "in_1.doc"))


def test_word_automation_call_sequence(tmp_path, monkeypatch):
    calls = []

    class FakeDoc:
        def SaveAs2(self, path, FileFormat):
            calls.append(("SaveAs2", os.path.basename(path), FileFormat))
            open(path, "wb").close()

        def Close(self, SaveChanges):
            calls.append(("Close", SaveChanges))

    class FakeDocuments:
        def Open(self, path, **kwargs):
            calls.append(("Open", os.path.basename(path), kwargs["ReadOnly"]))
            return FakeDoc()

    class FakeWord:
        Documents = FakeDocuments()

        def Quit(self):
            calls.append(("Quit", self.AutomationSecurity))

    client = types.SimpleNamespace(DispatchEx=lambda name: FakeWord())
    monkeypatch.setitem(sys.modules, "win32com", types.SimpleNamespace(client=client))
    monkeypatch.setitem(sys.modules, "win32com.client", client)
    monkeypatch.setitem(sys.modules, "pythoncom", types.SimpleNamespace(
        CoInitialize=lambda: calls.append(("CoInitialize",)),
        CoUninitialize=lambda: calls.append(("CoUninitialize",)),
    ))
    monkeypatch.setattr(legacy_convert.sys, "platform", "win32")

    out = convert_doc(str(tmp_path / "in_1.doc"))
    assert out == str(tmp_path / "in_1.docx")
    assert calls == [
        ("CoInitialize",),
        ("Open", "in_1.doc", True),
        ("SaveAs2", "in_1.docx", 12),
        ("Close", False),
        ("Quit", 3),  # macros were disabled
        ("CoUninitialize",),
    ]


@pytest.mark.nicegui_main_file("docflag.py")
async def test_doc_upload_is_converted_and_flagged(user: User, fake_soffice) -> None:
    await user.open("/")
    upload_element = next(iter(user.find(marker="doc-upload").elements))
    file_upload = Upload.SmallFileUpload(
        name="old report.doc", content_type="application/msword", _data=b"binary"
    )
    with user.client:
        await upload_element.handle_uploads([file_upload])
    await user.should_see("old report.doc")

    user.find(marker="flag-button").click()
    response = await user.download.next(timeout=10.0)
    doc = Document(io.BytesIO(response.content))
    assert [r.text for r in doc.paragraphs[0].runs if r.font.highlight_color] == ["draft"]
    await user.should_see("'old report_flagged.docx' is ready")


@pytest.mark.nicegui_main_file("docflag.py")
async def test_doc_upload_without_converter_is_refused(user: User, monkeypatch) -> None:
    monkeypatch.setattr(legacy_convert.sys, "platform", "linux")
    monkeypatch.setattr(legacy_convert, "_find_soffice", lambda: None)
    await user.open("/")
    upload_element = next(iter(user.find(marker="doc-upload").elements))
    file_upload = Upload.SmallFileUpload(
        name="old.doc", content_type="application/msword", _data=b"binary"
    )
    with user.client:
        await upload_element.handle_uploads([file_upload])
    await user.should_see("needs Microsoft Word or LibreOffice")
    await user.should_see("No documents added yet.")
