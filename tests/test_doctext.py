"""Word -> plain text for the comparison tab."""
from __future__ import annotations

import os
import sys

import pytest
from docx import Document
from lxml import etree

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docflag.compare.diffing import diff_texts
from docflag.doctext import extract_text
from docflag.package_io import Package
from docflag.xml_utils import NSMAP, qn
from tests.fixtures.make_fixtures import build_plain, build_sample

W = NSMAP["w"]
MC = NSMAP["mc"]


def _insert(path: str, xml: str, index: int = 0) -> None:
    pkg = Package.load(path)
    body = pkg.get_xml("word/document.xml").find(qn("w:body"))
    body.insert(index, etree.fromstring(xml))
    pkg.save(path)


def test_paragraphs_tables_and_no_headers(tmp_path):
    text = extract_text(build_sample(str(tmp_path / "s.docx")))
    lines = text.split("\n")
    assert lines[0] == "This report is Confidential and still a draft."
    assert lines[1] == "Marked as confidential by legal."  # split runs are joined
    assert "Status | DRAFT" in lines
    assert "Confidential header" not in text and "Plain footer" not in text


def test_table_rows_skip_empty_rows_and_flatten_nested_tables(tmp_path):
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Είδος"
    table.cell(0, 1).text = "Τιμή"
    table.cell(2, 0).text = "Καφές"
    table.cell(2, 1).add_paragraph("1,20")
    inner = table.cell(2, 1).add_table(rows=1, cols=2)
    inner.cell(0, 0).text = "x"
    inner.cell(0, 1).text = "a|b"
    path = str(tmp_path / "t.docx")
    doc.save(path)
    assert extract_text(path).split("\n") == ["Είδος | Τιμή", "Καφές | 1,20 x a b"]


def test_tracked_deletions_are_left_out(tmp_path):
    path = build_plain(str(tmp_path / "a.docx"), "after")
    _insert(path, (
        f'<w:p xmlns:w="{W}">'
        '<w:del w:id="1" w:author="x"><w:r><w:delText>gone</w:delText></w:r></w:del>'
        '<w:ins w:id="2" w:author="x"><w:r><w:t>kept</w:t></w:r></w:ins>'
        '</w:p>'
    ))
    assert extract_text(path).split("\n") == ["kept", "after"]


def test_text_box_is_read_once(tmp_path):
    path = build_plain(str(tmp_path / "a.docx"), "after")
    box = '<w:txbxContent><w:p><w:r><w:t>in the box</w:t></w:r></w:p></w:txbxContent>'
    _insert(path, (
        f'<w:p xmlns:w="{W}" xmlns:mc="{MC}"><w:r><w:t>anchor</w:t></w:r><w:r>'
        f'<mc:AlternateContent><mc:Choice Requires="wps"><w:drawing>{box}</w:drawing></mc:Choice>'
        f'<mc:Fallback><w:pict>{box}</w:pict></mc:Fallback></mc:AlternateContent>'
        '</w:r></w:p>'
    ))
    assert extract_text(path).split("\n") == ["anchor", "in the box", "after"]


def test_line_breaks_split_lines(tmp_path):
    path = build_plain(str(tmp_path / "a.docx"), "x")
    _insert(path, f'<w:p xmlns:w="{W}"><w:r><w:t>one</w:t><w:br/><w:t>two</w:t></w:r></w:p>')
    assert extract_text(path).split("\n")[:2] == ["one", "two"]


def test_changed_table_cell_is_diffed_row_by_row(tmp_path):
    def build(name, price):
        doc = Document()
        doc.add_paragraph("Τιμοκατάλογος")
        table = doc.add_table(rows=2, cols=2)
        for (r, c), value in {(0, 0): "Είδος", (0, 1): "Τιμή",
                              (1, 0): "Καφές", (1, 1): price}.items():
            table.cell(r, c).text = value
        path = str(tmp_path / name)
        doc.save(path)
        return path

    result = diff_texts(extract_text(build("a.docx", "1,20")), extract_text(build("b.docx", "1,50")))
    assert result.structured
    assert (result.stats.changed, result.stats.added, result.stats.removed) == (1, 0, 0)


def test_bad_file_is_rejected(tmp_path):
    path = tmp_path / "bad.docx"
    path.write_bytes(b"not a zip")
    with pytest.raises(ValueError):
        extract_text(str(path))
