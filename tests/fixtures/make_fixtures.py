"""Builds the synthetic .docx files the tests run against (via python-docx)."""
from __future__ import annotations

import os

from docx import Document
from docx.enum.text import WD_COLOR_INDEX


def build_sample(path: str) -> str:
    """One document exercising every place a word can hide."""
    doc = Document()
    doc.add_paragraph("This report is Confidential and still a draft.")

    # One visible word spread over two differently formatted runs.
    split = doc.add_paragraph("Marked as ")
    split.add_run("confi")
    split.add_run("dential").bold = True
    split.add_run(" by legal.")

    doc.add_paragraph("A draftsman redrafted it.")  # substrings only: no match
    doc.add_paragraph("For internal  use only, TBD.")  # phrase, double space

    already = doc.add_paragraph()
    already.add_run("draft").font.highlight_color = WD_COLOR_INDEX.GREEN

    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Status"
    table.cell(0, 1).text = "DRAFT"

    section = doc.sections[0]
    section.header.paragraphs[0].text = "Confidential header"
    section.footer.paragraphs[0].text = "Plain footer"

    os.makedirs(os.path.dirname(path), exist_ok=True)
    doc.save(path)
    return path


def build_plain(path: str, text: str) -> str:
    doc = Document()
    doc.add_paragraph(text)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    doc.save(path)
    return path
