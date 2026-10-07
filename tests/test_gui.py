"""GUI tests through NiceGUI's in-process client simulator (no browser)."""
from __future__ import annotations

import csv
import io
import os
import sys
import zipfile

import pytest
from docx import Document
from nicegui.elements.upload import Upload
from nicegui.testing import User

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docflag import replacements, wordlist
from tests.fixtures.make_fixtures import build_plain

DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


async def _upload(user: User, marker: str, name: str, data: bytes,
                  content_type: str = DOCX_CONTENT_TYPE) -> None:
    upload_element = next(iter(user.find(marker=marker).elements))
    file_upload = Upload.SmallFileUpload(name=name, content_type=content_type, _data=data)
    with user.client:
        await upload_element.handle_uploads([file_upload])


async def _upload_doc(user: User, path: str) -> None:
    with open(path, "rb") as f:
        await _upload(user, "doc-upload", os.path.basename(path), f.read())


def _highlighted(data: bytes) -> list[str]:
    doc = Document(io.BytesIO(data))
    return [
        r.text for p in doc.paragraphs for r in p.runs
        if r.font.highlight_color is not None
    ]


@pytest.mark.nicegui_main_file("docflag.py")
async def test_page_renders_expected_controls(user: User) -> None:
    await user.open("/")
    await user.should_see("DocFlag")
    await user.should_see("1. Προσθήκη εγγράφων")
    await user.should_see("2. Λέξεις προς επισήμανση")
    await user.should_see("3. Αντικαταστάσεις")
    await user.should_see("4. Επισήμανση")
    await user.should_see("τηλ. → τηλέφωνο")  # default replacements
    await user.should_see("Δεν έχουν προστεθεί έγγραφα ακόμη.")
    await user.should_see("confidential")  # default list
    await user.should_see("Άνοιγμα από άλλους υπολογιστές του δικτύου")


@pytest.mark.nicegui_main_file("docflag.py")
async def test_flag_button_warns_with_no_documents(user: User) -> None:
    await user.open("/")
    user.find(marker="flag-button").click()
    await user.should_see("Προσθέστε πρώτα τουλάχιστον ένα αρχείο Word.")


@pytest.mark.nicegui_main_file("docflag.py")
async def test_non_word_upload_is_skipped(user: User) -> None:
    await user.open("/")
    await _upload(user, "doc-upload", "notes.pdf", b"x", "application/pdf")
    await user.should_see("Δεν έχουν προστεθεί έγγραφα ακόμη.")


@pytest.mark.nicegui_main_file("docflag.py")
async def test_add_word_then_flag_single_document(user: User, tmp_path) -> None:
    path = build_plain(str(tmp_path / "memo.docx"), "The zebra is a draft animal.")
    await user.open("/")

    user.find(marker="new-word").type("Zebra")
    user.find(marker="add-word").click()
    await user.should_see("Zebra")
    assert "Zebra" in wordlist.load()

    await _upload_doc(user, path)
    await user.should_see("memo.docx")

    user.find(marker="flag-button").click()
    response = await user.download.next(timeout=10.0)
    assert response.status_code == 200
    assert _highlighted(response.content) == ["zebra", "draft"]
    await user.should_see("2 εμφανίσεις")
    await user.should_see("Το 'memo_flagged.docx' είναι έτοιμο")


@pytest.mark.nicegui_main_file("docflag.py")
async def test_several_documents_come_back_as_zip_with_report(user: User, tmp_path) -> None:
    wordlist.save(["draft"])
    replacements.save([("another", "one more")])
    first = build_plain(str(tmp_path / "one" / "memo.docx"), "A draft, another draft.")
    second = build_plain(str(tmp_path / "two" / "memo.docx"), "Nothing to see.")
    await user.open("/")
    await _upload_doc(user, first)
    await _upload_doc(user, second)

    user.find(marker="flag-button").click()
    response = await user.download.next(timeout=10.0)
    with zipfile.ZipFile(io.BytesIO(response.content)) as zf:
        assert sorted(zf.namelist()) == [
            "memo_flagged (2).docx", "memo_flagged.docx", "report.csv",
        ]
        assert _highlighted(zf.read("memo_flagged.docx")) == ["draft", "one more", "draft"]
        assert _highlighted(zf.read("memo_flagged (2).docx")) == []
        rows = list(csv.reader(io.StringIO(zf.read("report.csv").decode("utf-8-sig"))))
    assert rows == [
        ["Αρχείο", "Ενέργεια", "Λέξη", "Εμφανίσεις"],
        ["memo.docx", "Αντικατάσταση", "another → one more", "1"],
        ["memo.docx", "Επισήμανση", "draft", "2"],
        ["memo.docx", "", "(καμία εμφάνιση)", "0"],
    ]
    await user.should_see("Καμία εμφάνιση")


@pytest.mark.nicegui_main_file("docflag.py")
async def test_bulk_add_import_and_reset(user: User) -> None:
    wordlist.save(["one"])
    await user.open("/")

    user.find(marker="bulk-words").type("two\nthree words\nONE")
    user.find(marker="bulk-add").click()
    await user.should_see("three words")
    assert wordlist.load() == ["one", "two", "three words"]

    await _upload(user, "import-words", "list.txt", "four\n# skip\n".encode(), "text/plain")
    await user.should_see("four")
    assert wordlist.load() == ["one", "two", "three words", "four"]

    user.find(marker="reset-words").click()
    await user.should_see("Να αντικατασταθεί η κοινόχρηστη λίστα λέξεων με την προεπιλεγμένη;")
    user.find(marker="confirm-yes").click()
    await user.should_see("confidential")
    assert wordlist.load() == wordlist.defaults()


@pytest.mark.nicegui_main_file("docflag.py")
async def test_word_with_a_leading_space(user: User) -> None:
    wordlist.save([])
    await user.open("/")
    user.find(marker="new-word").type(" .")
    user.find(marker="add-word").click()
    await user.should_see("␣.")
    assert wordlist.load() == [" ."]


@pytest.mark.nicegui_main_file("docflag.py")
async def test_add_replacement_then_process(user: User, tmp_path) -> None:
    wordlist.save([])
    replacements.save([])
    path = build_plain(str(tmp_path / "memo.docx"), "Τηλ. 210 .")
    await user.open("/")

    user.find(marker="replace-from").type("τηλ.")
    user.find(marker="replace-to").type("τηλέφωνο")
    user.find(marker="add-replacement").click()
    await user.should_see("τηλ. → τηλέφωνο")
    user.find(marker="bulk-replacements").type('" ." => "."\nno separator')
    user.find(marker="bulk-add-replacements").click()
    await user.should_see("␣. → .")
    await user.should_see("1 γραμμή παραλείφθηκε")
    assert replacements.load() == [("τηλ.", "τηλέφωνο"), (" .", ".")]

    await _upload_doc(user, path)
    user.find(marker="flag-button").click()
    response = await user.download.next(timeout=10.0)
    doc = Document(io.BytesIO(response.content))
    assert doc.paragraphs[0].text == "Τηλέφωνο 210."
    assert _highlighted(response.content) == ["Τηλέφωνο", "."]
    await user.should_see("2 αντικαταστάσεις")


@pytest.mark.nicegui_main_file("docflag.py")
async def test_replacement_import_and_reset(user: User) -> None:
    replacements.save([])
    await user.open("/")
    await _upload(
        user, "import-replacements", "list.txt",
        "οδ. => οδός\n# skip\n".encode(), "text/plain",
    )
    await user.should_see("οδ. → οδός")
    assert replacements.load() == [("οδ.", "οδός")]

    user.find(marker="reset-replacements").click()
    await user.should_see("Να αντικατασταθεί η κοινόχρηστη λίστα αντικαταστάσεων με την προεπιλεγμένη;")
    user.find(marker="confirm-yes").click()
    await user.should_see("τηλ. → τηλέφωνο")
    assert replacements.load() == replacements.defaults()
