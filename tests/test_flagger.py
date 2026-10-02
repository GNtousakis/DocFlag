"""Engine tests: what gets highlighted, what doesn't, and that the text of
the document is never altered."""
from __future__ import annotations

import os
import sys

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from lxml import etree

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docflag.flagger import build_pattern, flag_document
from docflag.package_io import Package
from docflag.xml_utils import NSMAP, qn
from tests.fixtures.make_fixtures import build_plain, build_sample

WORDS = ["confidential", "draft", "internal use only", "TBD"]


def _highlighted(paragraph) -> list[str]:
    return [r.text for r in paragraph.runs if r.font.highlight_color is not None]


@pytest.fixture
def flagged(tmp_path):
    src = build_sample(str(tmp_path / "sample.docx"))
    dst = str(tmp_path / "sample_flagged.docx")
    result = flag_document(src, dst, WORDS)
    return Document(src), Document(dst), result


def test_text_is_unchanged(flagged):
    original, output, _ = flagged
    assert [p.text for p in output.paragraphs] == [p.text for p in original.paragraphs]


def test_counts(flagged):
    _, _, result = flagged
    assert result.counts == {
        "confidential": 3,  # body, split across runs, header
        "draft": 3,         # body, already-highlighted run, table
        "internal use only": 1,
        "TBD": 1,
    }
    assert result.total == 8


def test_case_insensitive_and_only_the_match_is_highlighted(flagged):
    _, output, _ = flagged
    assert _highlighted(output.paragraphs[0]) == ["Confidential", "draft"]


def test_match_split_across_runs_keeps_formatting(flagged):
    _, output, _ = flagged
    paragraph = output.paragraphs[1]
    assert _highlighted(paragraph) == ["confi", "dential"]
    bold = {r.text: bool(r.bold) for r in paragraph.runs}
    assert bold["confi"] is False and bold["dential"] is True


def test_whole_words_only(flagged):
    _, output, _ = flagged
    assert _highlighted(output.paragraphs[2]) == []


def test_phrase_matches_across_irregular_whitespace(flagged):
    _, output, _ = flagged
    assert _highlighted(output.paragraphs[3]) == ["internal  use only", "TBD"]


def test_existing_highlight_is_replaced(flagged):
    _, output, _ = flagged
    run = output.paragraphs[4].runs[0]
    assert run.font.highlight_color == WD_COLOR_INDEX.YELLOW
    assert len(run._r.rPr.findall(qn("w:highlight"))) == 1


def test_table_and_header_are_covered(flagged):
    _, output, _ = flagged
    assert _highlighted(output.tables[0].cell(0, 1).paragraphs[0]) == ["DRAFT"]
    header = output.sections[0].header.paragraphs[0]
    assert _highlighted(header) == ["Confidential"]
    assert _highlighted(output.sections[0].footer.paragraphs[0]) == []


def test_colour_choice(tmp_path):
    src = build_plain(str(tmp_path / "a.docx"), "a draft")
    dst = str(tmp_path / "b.docx")
    flag_document(src, dst, ["draft"], colour="cyan")
    run = Document(dst).paragraphs[0].runs[1]
    assert run.font.highlight_color == WD_COLOR_INDEX.TURQUOISE


def test_highlight_is_placed_in_schema_order(tmp_path):
    src = str(tmp_path / "a.docx")
    doc = Document()
    run = doc.add_paragraph().add_run("draft")
    run.bold = True
    run.underline = True
    doc.save(src)
    dst = str(tmp_path / "b.docx")
    flag_document(src, dst, ["draft"])
    rpr = Document(dst).paragraphs[0].runs[0]._r.rPr
    tags = [etree.QName(child).localname for child in rpr]
    assert tags.index("b") < tags.index("highlight") < tags.index("u")


def test_tracked_deletions_are_skipped_and_hyperlinks_covered(tmp_path):
    src = build_plain(str(tmp_path / "a.docx"), "placeholder")
    pkg = Package.load(src)
    w = NSMAP["w"]
    body = pkg.get_xml("word/document.xml").find(qn("w:body"))
    body.insert(0, etree.fromstring(
        f'<w:p xmlns:w="{w}">'
        '<w:del w:id="1" w:author="x"><w:r><w:delText>draft</w:delText></w:r></w:del>'
        '<w:hyperlink w:anchor="top"><w:r><w:t>the draft link</w:t></w:r></w:hyperlink>'
        '</w:p>'
    ))
    pkg.save(src)

    dst = str(tmp_path / "b.docx")
    result = flag_document(src, dst, ["draft"])
    assert result.counts == {"draft": 1}
    paragraph = Package.load(dst).get_xml("word/document.xml").find(f".//{qn('w:p')}")
    highlighted = [
        "".join(r.itertext()) for r in paragraph.iter(qn("w:r"))
        if r.find(f"{qn('w:rPr')}/{qn('w:highlight')}") is not None
    ]
    assert highlighted == ["draft"]
    assert paragraph.find(f".//{qn('w:delText')}").text == "draft"


def test_non_latin_words(tmp_path):
    src = build_plain(str(tmp_path / "a.docx"), "Το έγγραφο είναι ΕΜΠΙΣΤΕΥΤΙΚΌ.")
    dst = str(tmp_path / "b.docx")
    result = flag_document(src, dst, ["εμπιστευτικό"])
    assert result.counts == {"εμπιστευτικό": 1}


def _flag_text(tmp_path, text, words):
    src = build_plain(str(tmp_path / "a.docx"), text)
    dst = str(tmp_path / "b.docx")
    result = flag_document(src, dst, words)
    output = Document(dst)
    assert output.paragraphs[0].text == text
    return result.counts, _highlighted(output.paragraphs[0])


def test_accents_are_ignored_both_ways(tmp_path):
    counts, marked = _flag_text(
        tmp_path, "ΑΥΣΤΗΡΑ ΕΜΠΙΣΤΕΥΤΙΚΟ, εμπιστευτικό, café", ["εμπιστευτικό", "cafe"]
    )
    assert counts == {"εμπιστευτικό": 2, "cafe": 1}
    assert marked == ["ΕΜΠΙΣΤΕΥΤΙΚΟ", "εμπιστευτικό", "café"]


def test_decomposed_accents_are_highlighted_whole(tmp_path):
    text = "ένα προσχε\u0301διο\u0301 εδώ"  # accents as separate characters
    counts, marked = _flag_text(tmp_path, text, ["προσχεδιο"])
    assert counts == {"προσχεδιο": 1}
    assert marked == ["προσχε\u0301διο\u0301"]


def test_wildcard_matches_endings_but_stays_within_the_word(tmp_path):
    counts, marked = _flag_text(
        tmp_path,
        "του προσχεδίου, τα ΠΡΟΣΧΕΔΙΑ, απροσχεδίαστο, προσχέδιο τελικό",
        ["προσχέδι*", "*"],
    )
    assert counts == {"προσχέδι*": 3}
    assert marked == ["προσχεδίου", "ΠΡΟΣΧΕΔΙΑ", "προσχέδιο"]


def test_final_sigma(tmp_path):
    counts, _ = _flag_text(tmp_path, "Ο ΛΟΓΟΣ και ο λόγος", ["λόγος"])
    assert counts == {"λόγος": 2}


def test_longest_phrase_wins():
    pattern = build_pattern(["internal", "internal use only"])
    assert pattern.search("for internal use only").group() == "internal use only"


def test_empty_list_and_bad_file_are_rejected(tmp_path):
    src = build_plain(str(tmp_path / "a.docx"), "text")
    with pytest.raises(ValueError):
        flag_document(src, str(tmp_path / "b.docx"), ["  "])
    bogus = tmp_path / "bogus.docx"
    bogus.write_bytes(b"not a zip")
    with pytest.raises(ValueError):
        flag_document(str(bogus), str(tmp_path / "c.docx"), ["draft"])
