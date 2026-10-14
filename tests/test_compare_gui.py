"""Σύγκριση tab tests through NiceGUI's in-process client simulator, with
the stub OCR engine (see conftest.py)."""
from __future__ import annotations

import asyncio
import io
import os
import sys

import pytest
from docx import Document
from nicegui.elements.upload import Upload
from nicegui.testing import User
from PIL import Image, ImageDraw, ImageFont

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docflag.pdftext import minimal_pdf
from tests.test_gui import DOCX_CONTENT_TYPE, _upload

pytestmark = pytest.mark.nicegui_main_file("docflag.py")


def _docx(paragraph: str, price: str) -> bytes:
    doc = Document()
    doc.add_paragraph(paragraph)
    table = doc.add_table(rows=2, cols=2)
    for (r, c), value in {(0, 0): "Είδος", (0, 1): "Τιμή", (1, 0): "Καφές", (1, 1): price}.items():
        table.cell(r, c).text = value
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _png(last_line: str) -> bytes:
    img = Image.new("RGB", (900, 300), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.load_default(size=32)
    for i, line in enumerate(["Quarterly report invoice total", "Payment delivery agreement", last_line]):
        draw.text((30, 20 + i * 60), line, fill="black", font=font)
    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    return buffer.getvalue()


def _textarea(user: User, side: str):
    return next(iter(user.find(marker=f"text-{side}").elements))


async def _wait_for(condition, timeout: float = 5.0) -> None:
    """Uploads are handled in background tasks: wait for their effect."""
    for _ in range(int(timeout / 0.05)):
        if condition():
            return
        await asyncio.sleep(0.05)
    raise AssertionError("timed out")


async def _open(user: User) -> None:
    await user.open("/")
    user.find(marker="compare-tab").click()


async def test_compare_tab_renders(user: User) -> None:
    await _open(user)
    await user.should_see("Έγγραφο Α")
    await user.should_see("Έγγραφο Β")
    await user.should_see(marker="compare-button")
    await user.should_see("Δοκιμαστική λειτουργία")  # no Windows OCR here


async def test_compare_needs_text_on_both_sides(user: User) -> None:
    await _open(user)
    user.find(marker="compare-button").click()
    await user.should_see("Και τα δύο έγγραφα πρέπει να έχουν κείμενο.")


async def test_two_word_files(user: User) -> None:
    await _open(user)
    await _upload(user, "upload-A", "a.docx", _docx("Τιμοκατάλογος", "1,20"))
    await _upload(user, "upload-B", "b.docx", _docx("Τιμοκατάλογος", "1,50"))
    await _wait_for(lambda: _textarea(user, "A").value and _textarea(user, "B").value)
    assert _textarea(user, "A").value == "Τιμοκατάλογος\nΕίδος | Τιμή\nΚαφές | 1,20"
    await user.should_see("Κείμενο από το 'b.docx'.")

    user.find(marker="compare-button").click()
    await user.should_see("Αλλαγές: 1")
    await user.should_see("Προσθήκες: 0")
    await user.should_see(marker="similarity")
    user.find(marker="compare-back").click()
    await user.should_see(marker="compare-button")


async def test_screenshots_are_read_and_compared(user: User) -> None:
    await _open(user)
    await _upload(user, "upload-A", "image.png", _png("x"), "image/png")
    await _wait_for(lambda: _textarea(user, "A").value)
    await _upload(user, "upload-B", "image.png", _png("y"), "image/png")
    await _wait_for(lambda: _textarea(user, "B").value)
    assert "Μήλα | 10 | 1,20" in _textarea(user, "A").value  # stub: table_v1
    assert "Μήλα | 12 | 1,20" in _textarea(user, "B").value  # stub: table_v2

    user.find(marker="compare-button").click()
    await user.should_see("Αλλαγές: 2")
    await user.should_see("Διαγραφές: 0")


async def test_appended_screenshot_text_can_be_removed(user: User) -> None:
    await _open(user)
    await _upload(user, "upload-A", "one.png", _png("x"), "image/png")
    await _upload(user, "upload-A", "two.png", _png("x"), "image/png")
    await _wait_for(lambda: "Πορτοκάλια | 8 | 1,80" in (_textarea(user, "A").value or ""))
    first, second = _textarea(user, "A").value.split("\nΕίδος", 1)
    assert "1,50" in first and "1,80" in second  # read in order

    user.find(marker="remove-region-A-2").click()
    await _wait_for(lambda: "1,80" not in _textarea(user, "A").value)
    assert _textarea(user, "A").value == first


async def test_word_file_is_refused_in_image_mode(user: User) -> None:
    await _open(user)
    next(iter(user.find(marker="compare-mode").elements)).value = "image"
    await _upload(user, "upload-A", "a.docx", _docx("x", "1"), DOCX_CONTENT_TYPE)
    await user.should_see("Στη σύγκριση εικόνων προσθέστε εικόνες, όχι αρχεία Word.")


async def test_image_mode(user: User) -> None:
    await _open(user)
    next(iter(user.find(marker="compare-mode").elements)).value = "image"
    await _upload(user, "upload-A", "a.png", _png("Customer reference number 88231"), "image/png")
    await _upload(user, "upload-B", "b.png", _png("Customer reference number 99999"), "image/png")
    await user.should_see("900×300 px")
    await asyncio.sleep(0.2)
    user.find(marker="compare-button").click()
    await user.should_see(marker="count-regions")
    label = next(iter(user.find(marker="count-regions").elements))
    assert label.text != "Περιοχές διαφορών: 0"


def _pdf(price: str, pages: int = 1) -> bytes:
    table = [(72, 80, "Price list"), (72, 120, "Item"), (232, 120, "Price"),
             (72, 142, "Apples"), (232, 142, price)]
    return minimal_pdf([table] + [[(72, 80, "Prices include VAT.")]] * (pages - 1))


async def test_two_pdf_files(user: User) -> None:
    await _open(user)
    await _upload(user, "upload-A", "a.pdf", _pdf("1,20"), "application/pdf")
    await _upload(user, "upload-B", "b.pdf", _pdf("1,50", pages=2), "application/pdf")
    await _wait_for(lambda: _textarea(user, "A").value and _textarea(user, "B").value)
    assert _textarea(user, "A").value == "Price list\nItem | Price\nApples | 1,20"
    await user.should_see("Κείμενο από το 'b.pdf' (2 σελίδες).")

    user.find(marker="compare-button").click()
    await user.should_see("Αλλαγές: 1")
    await user.should_see("Προσθήκες: 1")


async def test_pdf_in_image_mode_uses_the_first_page(user: User) -> None:
    await _open(user)
    next(iter(user.find(marker="compare-mode").elements)).value = "image"
    await _upload(user, "upload-A", "a.pdf", _pdf("1,20", pages=3), "application/pdf")
    await user.should_see("Η πρώτη σελίδα του 'a.pdf' (3 σελίδες).")
